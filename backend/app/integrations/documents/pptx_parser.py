from pptx import Presentation
from app.integrations.documents.base import ParsedPage


class PptxParser:
    def parse_file(self, file_path: str) -> list[ParsedPage]:
        prs = Presentation(file_path)
        pages: list[ParsedPage] = []

        for idx, slide in enumerate(prs.slides, start=1):
            slide_texts = []
            slide_title = None
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        txt = paragraph.text.strip()
                        if txt:
                            slide_texts.append(txt)
            if slide.shapes.title and slide.shapes.title.has_text_frame:
                slide_title = slide.shapes.title.text_frame.text.strip()

            if slide_texts:
                pages.append(ParsedPage(
                    page_number=idx,
                    text="\n\n".join(slide_texts),
                    heading=slide_title or f"Slide {idx}",
                ))
        return pages
